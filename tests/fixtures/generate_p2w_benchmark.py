"""Synthetic corpus and ground truth for the M5B PDF to Word benchmark.

Writes ignored data to .tools/p2w-benchmarks/{corpus,truth}. Every file is
authored here from generated text; no user documents are involved. The corpus
is deterministic (seeded). Ground truth is produced by this script, never by a
candidate converter.

Criteria are fixed in docs/benchmarks/M5B-pdf-to-word-criteria.md.

    uv run --with python-docx --with reportlab --with pypdfium2 --with pillow \
        --with numpy python tests/fixtures/generate_p2w_benchmark.py

Producers: LibreOffice (python-docx -> soffice), Chromium (HTML -> PDF) and
ReportLab (direct PDF), so results are not tied to one PDF writer.

Truth schema (JSON):
  blocks   reading-order list of {type, text} or {type: "table", rows: [[cell]]}
  images   embedded raster images in the source
  furniture  words that repeat per page (running header/footer); order-exempt
  scanned  True when the PDF has no text layer
  hidden   words visible only as pixels (scans): must not appear in a DOCX
"""

import argparse
import json
import random
import shutil
import subprocess
import tempfile
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from PIL import Image, ImageDraw
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[2] / ".tools" / "p2w-benchmarks"
CORPUS = ROOT / "corpus"
TRUTH = ROOT / "truth"
CHROMIUM = next(Path("/opt/pw-browsers").glob("chromium-*/chrome-linux/chrome"), None)
SYLLABLES = "ka ro mi tel von dar lin sor pex bu qui nor fa gel hap jus lor mez nid ost ver tam dru yel zan".split()
ARABIC = (
    "الكتاب المدرسة الطالب المعلم البيت الشارع المدينة الصباح المساء النهار "
    "اللغة العمل الطريق الماء الشمس القمر السماء الأرض الحديقة المكتبة الرسالة"
).split()


class Corpus:
    def __init__(self, seed=20261021):
        self.rng = random.Random(seed)
        self.vocab = []
        seen = set()
        while len(self.vocab) < 900:
            word = "".join(self.rng.choice(SYLLABLES) for _ in range(self.rng.choice((2, 3))))
            if word not in seen:
                seen.add(word)
                self.vocab.append(word)

    def words(self, n):
        return [self.rng.choice(self.vocab) for _ in range(n)]

    def sentence(self):
        w = self.words(self.rng.randint(8, 15))
        w[0] = w[0].capitalize()
        return " ".join(w) + "."

    def paragraph(self, sentences=None):
        return " ".join(self.sentence() for _ in range(sentences or self.rng.randint(4, 7)))

    def title(self, n=4):
        return " ".join(w.capitalize() for w in self.words(n))

    def cell(self):
        return " ".join(self.words(self.rng.choice((1, 1, 2))))


def para(kind, text):
    return {"type": kind, "text": text}


def soffice_pdf(docx_path, out_dir):
    profile = tempfile.mkdtemp(prefix="lo-profile-")
    try:
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            check=True, capture_output=True, timeout=300,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def chromium_pdf(html, target):
    if CHROMIUM is None:
        raise SystemExit("Chromium not found under /opt/pw-browsers")
    with tempfile.TemporaryDirectory() as tmp:
        page = Path(tmp) / "page.html"
        page.write_text(html, encoding="utf-8")
        subprocess.run(
            [str(CHROMIUM), "--headless=new", "--no-sandbox", "--disable-gpu",
             "--no-pdf-header-footer", f"--print-to-pdf={target}", f"file://{page}"],
            check=True, capture_output=True, timeout=180,
        )


CSS = """
@page { size: A4; margin: 18mm; %s }
body { font-family: 'Liberation Serif', serif; font-size: 11pt; line-height: 1.35; }
h1 { font-size: 22pt; margin: 0 0 8pt; } h2 { font-size: 15pt; margin: 14pt 0 4pt; }
p { margin: 0 0 8pt; } table { border-collapse: collapse; width: 100%%; margin: 6pt 0 10pt; }
td, th { border: 1px solid #444; padding: 3pt 5pt; text-align: left; font-size: 10pt; }
th { background: #e6e6e6; } thead { display: table-row-group; }
figure { margin: 8pt 0; } figcaption { font-size: 9pt; font-style: italic; }
"""


def page(body, extra_css="", margin_boxes="", lang="en", direction="ltr", family=None):
    font = f"body {{ font-family: {family}; }}" if family else ""
    return (
        f"<!doctype html><html lang='{lang}' dir='{direction}'><head><meta charset='utf-8'>"
        f"<style>{CSS % margin_boxes}{font}{extra_css}</style></head><body>{body}</body></html>"
    )


def make_image(path, seed, size=(600, 400)):
    r = random.Random(seed)
    img = Image.new("RGB", size, (r.randint(60, 200), r.randint(60, 200), r.randint(60, 200)))
    d = ImageDraw.Draw(img)
    for _ in range(14):
        x0, y0 = r.randint(0, size[0] - 50), r.randint(0, size[1] - 50)
        d.ellipse([x0, y0, x0 + r.randint(40, 200), y0 + r.randint(40, 160)],
                  fill=(r.randint(0, 255), r.randint(0, 255), r.randint(0, 255)))
    img.save(path, "JPEG", quality=88)


# ---- file builders: each returns (blocks, extra) and writes CORPUS/<id>.pdf ----


def build_a01(c, out):
    doc = Document()
    blocks = []
    title = c.title()
    doc.add_heading(title, level=1)
    blocks.append(para("heading", title))
    for _ in range(5):
        heading = c.title(3)
        doc.add_heading(heading, level=2)
        blocks.append(para("heading", heading))
        for _ in range(3):
            text = c.paragraph()
            doc.add_paragraph(text)
            blocks.append(para("paragraph", text))
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "A01.docx"
        doc.save(src)
        soffice_pdf(src, Path(tmp))
        shutil.copy(Path(tmp) / "A01.pdf", out)
    return blocks, {}


def build_a02(c, out):
    blocks, html = [], []
    title = c.title()
    html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    for section in range(3):
        heading = c.title(3)
        html.append(f"<h2>{heading}</h2>"); blocks.append(para("heading", heading))
        first = c.words(6); second = c.words(4); rest = c.paragraph(3)
        text = f"{' '.join(first)} {' '.join(second)} {rest}"
        html.append(f"<p><b>{' '.join(first)}</b> <i>{' '.join(second)}</i> {rest}</p>")
        blocks.append(para("paragraph", text))
        tag = "ul" if section % 2 == 0 else "ol"
        items = [c.sentence() for _ in range(4)]
        html.append(f"<{tag}>" + "".join(f"<li>{i}</li>" for i in items) + f"</{tag}>")
        blocks.extend(para("list_item", i) for i in items)
        for _ in range(2):
            text = c.paragraph(); html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
    chromium_pdf(page("".join(html)), out)
    return blocks, {}


def build_a03(c, out):
    blocks, html = [], []
    header, footer = "hdrmark Quarterly review", "ftrmark Internal"
    title = c.title()
    html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    for _ in range(6):
        heading = c.title(3)
        html.append(f"<h2>{heading}</h2>"); blocks.append(para("heading", heading))
        for _ in range(3):
            text = c.paragraph(); html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
    boxes = (f"@top-center {{ content: '{header}'; font-size: 9pt; }} "
             f"@bottom-left {{ content: '{footer}'; font-size: 9pt; }} "
             "@bottom-right { content: 'Page ' counter(page); font-size: 9pt; }")
    chromium_pdf(page("".join(html), margin_boxes=boxes), out)
    pages = len(pdfium.PdfDocument(str(out)))
    furniture = []
    for number in range(1, pages + 1):
        furniture += header.split() + footer.split() + ["Page", str(number)]
    return blocks, {"furniture": furniture}


def table_html(rows, header=True):
    out = ["<table>"]
    for index, row in enumerate(rows):
        tag = "th" if header and index == 0 else "td"
        out.append("<tr>" + "".join(f"<{tag}>{cell}</{tag}>" for cell in row) + "</tr>")
    out.append("</table>")
    return "".join(out)


def build_a04(c, out):
    blocks, html = [], []
    company = c.title(2)
    html.append(f"<h1>{company}</h1>"); blocks.append(para("heading", company))
    address = c.paragraph(2); html.append(f"<p>{address}</p>"); blocks.append(para("paragraph", address))
    meta = [["Invoice", c.cell()], ["Date", c.cell()], ["Customer", c.cell()], ["Terms", c.cell()]]
    html.append(table_html(meta, header=False)); blocks.append({"type": "table", "rows": meta})
    items = [["Item", "Description", "Qty", "Unit", "Total"]]
    for _ in range(12):
        items.append([c.cell(), c.cell(), str(c.rng.randint(1, 40)), f"{c.rng.randint(5, 900)}.00", f"{c.rng.randint(50, 9000)}.00"])
    html.append(table_html(items)); blocks.append({"type": "table", "rows": items})
    note = c.paragraph(2); html.append(f"<p>{note}</p>"); blocks.append(para("paragraph", note))
    chromium_pdf(page("".join(html)), out)
    return blocks, {}


def build_a05(c, out):
    blocks, html = [], []
    title = c.title(); html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    rows = [[c.cell().capitalize() for _ in range(6)]]
    for _ in range(70):
        rows.append([c.cell() for _ in range(4)] + [str(c.rng.randint(1, 9999)), f"{c.rng.randint(1, 99)}.{c.rng.randint(10, 99)}"])
    html.append(table_html(rows)); blocks.append({"type": "table", "rows": rows})
    chromium_pdf(page("".join(html)), out)
    return blocks, {}


def build_a06(c, out):
    blocks, html = [], []
    title = c.title(); html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    with tempfile.TemporaryDirectory() as tmp:
        for index in range(3):
            text = c.paragraph(); html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
            image = Path(tmp) / f"img{index}.jpg"; make_image(image, 100 + index)
            caption = "Figure " + str(index + 1) + " " + c.title(3)
            html.append(f"<figure><img src='file://{image}' width='360'><figcaption>{caption}</figcaption></figure>")
            blocks.append({"type": "image", "text": ""}); blocks.append(para("caption", caption))
        text = c.paragraph(); html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
        chromium_pdf(page("".join(html)), out)
    return blocks, {"images": 3}


def build_columns(c, out, count, sidebar=False):
    blocks, html = [], []
    title = c.title(); html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    if sidebar:
        side = [c.sentence() for _ in range(3)]
        html.append("<div style='float:left;width:22%;border:1px solid #444;padding:6pt;margin-right:8pt'>"
                    + "".join(f"<p>{s}</p>" for s in side) + "</div>")
        blocks.extend(para("sidebar", s) for s in side)
    paras = [c.paragraph(5) for _ in range(9 if count == 2 else 8)]
    html.append(f"<div style='column-count:{count};column-gap:16pt'>"
                + "".join(f"<p>{p}</p>" for p in paras) + "</div>")
    blocks.extend(para("paragraph", p) for p in paras)
    chromium_pdf(page("".join(html)), out)
    return blocks, {"columns": count}


def build_b03(c, out):
    pdf = canvas.Canvas(str(out), pagesize=landscape(A4))
    blocks = []
    width, height = landscape(A4)
    for _ in range(4):
        title = c.title(3)
        pdf.setFont("Helvetica-Bold", 30); pdf.drawString(60, height - 90, title)
        blocks.append(para("heading", title))
        for row in range(4):
            text = " ".join(c.words(c.rng.randint(5, 9)))
            pdf.setFont("Helvetica", 20); pdf.drawString(90, height - 170 - row * 55, text)
            blocks.append(para("bullet", text))
        pdf.showPage()
    pdf.save()
    return blocks, {}


def build_b04(c, out):
    pdf = canvas.Canvas(str(out), pagesize=A4)
    blocks = []
    width, height = A4
    title = c.title(3); pdf.setFont("Helvetica-Bold", 20); pdf.drawString(50, height - 70, title)
    blocks.append(para("heading", title))
    y = height - 130
    for _ in range(9):
        left, right = c.cell().capitalize() + ":", c.cell().capitalize() + ":"
        pdf.setFont("Helvetica", 11)
        pdf.drawString(50, y, left); pdf.line(130, y - 2, 270, y - 2)
        pdf.drawString(300, y, right); pdf.line(380, y - 2, 540, y - 2)
        blocks.extend([para("label", left), para("label", right)])
        y -= 34
    for _ in range(4):
        label = " ".join(c.words(3))
        pdf.rect(50, y - 2, 10, 10); pdf.drawString(68, y, label)
        blocks.append(para("label", label)); y -= 24
    pdf.showPage(); pdf.save()
    return blocks, {}


def build_b05(c, out):
    blocks, html = [], []
    title = c.title(); html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    with tempfile.TemporaryDirectory() as tmp:
        image = Path(tmp) / "wrap.jpg"; make_image(image, 777, (300, 220))
        html.append(f"<img src='file://{image}' width='220' style='float:left;margin:0 12pt 6pt 0'>")
        blocks.append({"type": "image", "text": ""})
        for _ in range(6):
            text = c.paragraph(5); html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
        chromium_pdf(page("".join(html)), out)
    return blocks, {"images": 1}


def build_b06(c, out):
    blocks, html = [], []
    title = " ".join(c.rng.choice(ARABIC) for _ in range(3))
    html.append(f"<h1>{title}</h1>"); blocks.append(para("heading", title))
    for _ in range(6):
        text = " ".join(c.rng.choice(ARABIC) for _ in range(c.rng.randint(40, 70)))
        html.append(f"<p>{text}</p>"); blocks.append(para("paragraph", text))
    chromium_pdf(page("".join(html), lang="ar", direction="rtl", family="'DejaVu Sans', sans-serif"), out)
    return blocks, {"rtl": True}


def scanned_pages(source, count=2, dpi=170):
    document = pdfium.PdfDocument(str(source))
    images = []
    for index in range(min(count, len(document))):
        images.append(document[index].render(scale=dpi / 72).to_pil().convert("RGB"))
    return images


def build_c01(c, out, source):
    images = scanned_pages(source)
    pdf = canvas.Canvas(str(out), pagesize=A4)
    for image in images:
        pdf.drawImage(ImageReader(image), 0, 0, width=A4[0], height=A4[1]); pdf.showPage()
    pdf.save()
    return [], {"scanned": True}


def build_c02(c, out, source, hidden_words):
    images = scanned_pages(source)
    pdf = canvas.Canvas(str(out), pagesize=A4)
    cursor = 0
    blocks = []
    for image in images:
        pdf.drawImage(ImageReader(image), 0, 0, width=A4[0], height=A4[1])
        text = pdf.beginText(50, A4[1] - 60); text.setTextRenderMode(3); text.setFont("Helvetica", 10)
        for _ in range(45):
            line = " ".join(hidden_words[cursor:cursor + 12]); cursor += 12
            text.textLine(line)
            if line: blocks.append(para("paragraph", line))
        pdf.drawText(text); pdf.showPage()
    pdf.save()
    return blocks, {"invisible_text_layer": True}


def build_long(c, out, pages_target):
    doc = Document()
    blocks = []
    words_per_page = 380
    count = 0
    while count < pages_target * words_per_page:
        heading = c.title(3); doc.add_heading(heading, level=2); blocks.append(para("heading", heading))
        for _ in range(4):
            text = c.paragraph(6); doc.add_paragraph(text); blocks.append(para("paragraph", text))
            count += len(text.split())
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "long.docx"; doc.save(src); soffice_pdf(src, Path(tmp))
        shutil.copy(Path(tmp) / "long.pdf", out)
    return blocks, {}


def build_d03(c, out):
    """About 19 MiB: noise-textured JPEG pages with a line of text each."""
    pdf = canvas.Canvas(str(out), pagesize=A4)
    blocks = []
    r = random.Random(5)
    with tempfile.TemporaryDirectory() as tmp:
        for index in range(8):
            pix = bytes(r.getrandbits(8) for _ in range(1700 * 1200 * 3))
            image = Image.frombytes("RGB", (1700, 1200), pix)
            image.save(Path(tmp) / f"n{index}.jpg", "JPEG", quality=92)
            pdf.drawImage(str(Path(tmp) / f"n{index}.jpg"), 36, 260, width=520, height=367)
            text = " ".join(c.words(10)); pdf.setFont("Helvetica", 12); pdf.drawString(50, 700, text)
            blocks.append(para("paragraph", text)); pdf.showPage()
    pdf.save()
    return blocks, {"images": 8}


def finish(file_id, klass, producer, out, blocks, extra):
    document = pdfium.PdfDocument(str(out))
    truth = {
        "id": file_id, "class": klass, "producer": producer,
        "pages": len(document), "bytes": out.stat().st_size,
        "blocks": blocks, "images": extra.pop("images", 0),
        "furniture": extra.pop("furniture", []),
        "scanned": extra.pop("scanned", False), **extra,
    }
    document.close()
    (TRUTH / f"{file_id}.json").write_text(json.dumps(truth, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"{file_id:28s} {producer:10s} pages={truth['pages']:<4d} {truth['bytes']:>10,d} B")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-long", action="store_true")
    args = parser.parse_args()
    shutil.rmtree(ROOT, ignore_errors=True)
    CORPUS.mkdir(parents=True); TRUTH.mkdir(parents=True)
    c = Corpus()

    plan = [
        ("A01-prose-headings", "A", "LibreOffice", lambda out: build_a01(c, out)),
        ("A02-emphasis-lists", "A", "Chromium", lambda out: build_a02(c, out)),
        ("A03-report-furniture", "A", "Chromium", lambda out: build_a03(c, out)),
        ("A04-invoice-table", "A", "Chromium", lambda out: build_a04(c, out)),
        ("A05-data-table", "A", "Chromium", lambda out: build_a05(c, out)),
        ("A06-images-captions", "A", "Chromium", lambda out: build_a06(c, out)),
        ("B01-two-column", "B", "Chromium", lambda out: build_columns(c, out, 2)),
        ("B02-three-column-sidebar", "B", "Chromium", lambda out: build_columns(c, out, 3, sidebar=True)),
        ("B03-slides", "B", "ReportLab", lambda out: build_b03(c, out)),
        ("B04-form", "B", "ReportLab", lambda out: build_b04(c, out)),
        ("B05-image-wrap", "B", "Chromium", lambda out: build_b05(c, out)),
        ("B06-arabic-rtl", "B", "Chromium", lambda out: build_b06(c, out)),
    ]
    for file_id, klass, producer, build in plan:
        out = CORPUS / f"{file_id}.pdf"
        blocks, extra = build(out)
        finish(file_id, klass, producer, out, blocks, extra)

    source = CORPUS / "A01-prose-headings.pdf"
    a01_words = [w for b in json.loads((TRUTH / "A01-prose-headings.json").read_text())["blocks"] for w in b["text"].split()]
    out = CORPUS / "C01-scanned.pdf"
    blocks, extra = build_c01(c, out, source)
    extra["hidden"] = a01_words[:300]
    finish("C01-scanned", "C", "ReportLab", out, blocks, extra)
    out = CORPUS / "C02-scanned-ocr-layer.pdf"
    blocks, extra = build_c02(c, out, source, a01_words)
    finish("C02-scanned-ocr-layer", "C", "ReportLab", out, blocks, extra)

    if not args.skip_long:
        for file_id, pages in (("D01-50-pages", 50), ("D02-200-pages", 200)):
            out = CORPUS / f"{file_id}.pdf"
            blocks, extra = build_long(c, out, pages)
            finish(file_id, "D", "LibreOffice", out, blocks, extra)
        out = CORPUS / "D03-near-20MiB.pdf"
        blocks, extra = build_d03(c, out)
        finish("D03-near-20MiB", "D", "ReportLab", out, blocks, extra)

    out = CORPUS / "D04-encrypted.pdf"
    subprocess.run(["qpdf", "--encrypt", "user-pw", "owner-pw", "256", "--", str(source), str(out)], check=True)
    truth = {"id": "D04-encrypted", "class": "D", "producer": "qpdf", "pages": 0,
             "bytes": out.stat().st_size, "blocks": [], "images": 0, "furniture": [],
             "scanned": False, "expect_failure": "encrypted"}
    (TRUTH / "D04-encrypted.json").write_text(json.dumps(truth, indent=1) + "\n")
    print(f"{'D04-encrypted':28s} {'qpdf':10s} {truth['bytes']:>10,d} B")
    out = CORPUS / "D05-corrupt.pdf"
    data = source.read_bytes()
    out.write_bytes(data[: int(len(data) * 0.6)])
    truth = {"id": "D05-corrupt", "class": "D", "producer": "truncated", "pages": 0,
             "bytes": out.stat().st_size, "blocks": [], "images": 0, "furniture": [],
             "scanned": False, "expect_failure": "corrupt"}
    (TRUTH / "D05-corrupt.json").write_text(json.dumps(truth, indent=1) + "\n")
    print(f"{'D05-corrupt':28s} {'truncated':10s} {truth['bytes']:>10,d} B")


if __name__ == "__main__":
    main()
