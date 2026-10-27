#!/usr/bin/env python3
"""Build the M7G device test pack: a zip of small, synthetic files to put on a
real phone or PC and run through every tool (docs/qa/M7G-device-checklist.md).

    python3 tests/fixtures/generate_device_pack.py OUTPUT_DIR

Writes OUTPUT_DIR/document-platform-device-test-pack.zip (and the folder it was
zipped from). Nothing here is a real document: every file is generated or is one
of the project's own synthetic fixtures in tests/fixtures/pdf. ReportLab, pypdf,
Pillow and python-docx are authoring tools only; none is a project dependency.
"""

from __future__ import annotations

import io
import random
import shutil
import sys
import zipfile
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

FIXTURES = Path(__file__).resolve().parent / "pdf"
PASSWORD = "synthetic-fixture-only"


def font(size: int) -> ImageFont.ImageFont:
    for name in ("DejaVuSans-Bold.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def numbered_pdf(path: Path, pages: int, label: str) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setTitle(label)
    for number in range(1, pages + 1):
        pdf.setFont("Helvetica-Bold", 120)
        pdf.drawCentredString(306, 430, str(number))
        pdf.setFont("Helvetica", 18)
        pdf.drawCentredString(306, 360, f"{label}: page {number} of {pages}")
        pdf.showPage()
    pdf.save()


def image_heavy_pdf(path: Path, pages: int) -> None:
    """Pages that are one large photo-like image each, so the file is heavy.

    The image is placed with a plain `q w 0 0 h x y cm /Im Do Q` content stream
    (not a ReportLab form wrapper), the layout Compress can analyse."""
    from pypdf import PdfWriter
    from pypdf.generic import (
        ArrayObject,
        DecodedStreamObject,
        DictionaryObject,
        NameObject,
        NumberObject,
        StreamObject,
    )

    rng = random.Random(7)
    writer = PdfWriter()
    width, height = 1700, 2200
    for number in range(1, pages + 1):
        # A smooth, colourful field with fine grain: it behaves like a photograph
        # (not like a screenshot or a flat graphic) and shrinks when recompressed.
        small = Image.frombytes(
            "RGB", (24, 32), bytes(rng.randrange(256) for _ in range(24 * 32 * 3))
        )
        image = small.resize((width, height), Image.BICUBIC)
        grain = Image.effect_noise((width, height), 60).convert("RGB")
        image = Image.blend(image, grain, 0.08)
        draw = ImageDraw.Draw(image)
        draw.text((120, 120), f"IMAGE {number}", fill="white", font=font(220))
        buffer = io.BytesIO()
        image.save(buffer, "JPEG", quality=92)

        picture = StreamObject()
        picture._data = buffer.getvalue()
        picture.update(
            {
                NameObject("/Type"): NameObject("/XObject"),
                NameObject("/Subtype"): NameObject("/Image"),
                NameObject("/Width"): NumberObject(width),
                NameObject("/Height"): NumberObject(height),
                NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                NameObject("/BitsPerComponent"): NumberObject(8),
                NameObject("/Filter"): NameObject("/DCTDecode"),
            }
        )
        content = DecodedStreamObject()
        content.set_data(b"q 612 0 0 792 0 0 cm /Im0 Do Q")
        page = writer.add_blank_page(612, 792)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/XObject"): DictionaryObject(
                    {NameObject("/Im0"): writer._add_object(picture)}
                )
            }
        )
        page[NameObject("/Contents")] = writer._add_object(content)
    writer.add_metadata({"/Title": "Image-heavy test"})
    with open(path, "wb") as handle:
        writer.write(handle)


def transparent_png(path: Path) -> None:
    image = Image.new("RGBA", (1200, 800), (0, 0, 0, 0))
    draw = ImageDraw.Draw(image)
    draw.ellipse((100, 80, 1100, 720), fill=(30, 110, 200, 255))
    draw.ellipse((300, 220, 900, 580), fill=(0, 0, 0, 0))  # transparent hole
    draw.text((370, 330), "PNG", fill=(60, 60, 60, 255), font=font(160))
    image.save(path, "PNG")


def camera_style_jpeg(path: Path) -> None:
    """Pixels stored sideways with EXIF orientation 6, as a phone held upright
    writes them. Shown correctly, the picture is a portrait with 'TOP' at the top
    and an arrow pointing up."""
    upright = Image.new("RGB", (1500, 2000), (235, 235, 225))
    draw = ImageDraw.Draw(upright)
    draw.rectangle((0, 0, 1500, 300), fill=(200, 60, 40))
    draw.text((470, 60), "TOP", fill="white", font=font(200))
    draw.polygon([(750, 450), (450, 900), (1050, 900)], fill=(30, 110, 200))
    draw.rectangle((680, 900, 820, 1500), fill=(30, 110, 200))
    draw.text((330, 1650), "upright", fill=(60, 60, 60), font=font(160))
    stored = upright.rotate(90, expand=True)  # 2000 x 1500, sideways
    exif = Image.Exif()
    exif[0x0112] = 6  # display by rotating 90 degrees clockwise
    exif[0x010F] = "Synthetic Camera"
    exif[0x0110] = "Test Model 1"
    stored.save(path, "JPEG", quality=88, exif=exif)


def word_document(path: Path) -> None:
    from docx import Document

    document = Document()
    document.core_properties.author = "Synthetic Author"
    section = document.sections[0]
    section.header.paragraphs[0].text = "Device test header"
    section.footer.paragraphs[0].text = "Device test footer"
    document.add_heading("Device test document", level=0)
    document.add_paragraph(
        "This file is generated for testing. It has headings, a list, a table "
        "and a page break, so the converted PDF can be checked at a glance."
    )
    document.add_heading("A list", level=1)
    for item in ("First item", "Second item", "Third item"):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("A table", level=1)
    table = document.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for row in range(3):
        for column in range(3):
            table.cell(row, column).text = f"R{row + 1}C{column + 1}"
    document.add_page_break()
    document.add_heading("Second page", level=1)
    document.add_paragraph("If you can read this on page 2, the page break worked.")
    document.save(path)


README = f"""DOCUMENT PLATFORM - DEVICE TEST PACK
All files are synthetic. None contains real or personal data.
Instructions: docs/qa/M7G-device-checklist.md in the repository.

FILE                              USE IT FOR / WHAT YOU SHOULD SEE
01-merge-a-2-pages.pdf            Merge with 02. Pages marked A1, A2.
02-merge-b-2-pages.pdf            Pages marked B1, B2 (turned sideways / upside down).
                                  Merge 01 + 02 -> 4 pages in the order A1 A2 B1 B2.
03-four-pages.pdf                 Organize, Split, Extract, Delete, Rotate, PDF to JPG,
                                  PDF to Text. Pages marked P1 to P4 (P4 is cropped and
                                  turned on purpose).
04-fifty-pages.pdf                A long list on a phone: Organize scrolling and moving.
                                  Big page numbers 1 to 50.
05-two-hundred-pages.pdf          The page limit. Organize or Split should accept it.
06-image-heavy-5-pages.pdf        Heavy pages (about 5.4 MB): PDF to JPG, Compress,
                                  Organize on a low-memory phone. Compress should report
                                  about 80% smaller. Big image per page, numbered.
07-metadata-rich.pdf              Remove PDF Metadata. Has a fake title, author and more.
08-scanned-no-text.pdf            PDF to Text must say there is no text layer.
09-partly-scanned.pdf             PDF to Text: pages 1 and 3 have text, page 2 does not.
10-password-protected.pdf         Every PDF tool must refuse it, saying it is password
                                  protected (some browsers say "damaged or unsupported").
                                  Its password is {PASSWORD}; you never need it.
11-damaged.pdf                    Every PDF tool must refuse it as damaged or unsupported.
12-pdf-without-extension          A real PDF whose name has no .pdf. See the checklist.
13-transparent.png                Images to PDF with a PNG that has transparency: a blue
                                  ring with a see-through middle. In the PDF the middle
                                  must be white, never black.
14-camera-style-rotated.jpg       Images to PDF: must come out upright (TOP at the top,
                                  arrow up), not sideways.
15-word-document.docx             Word to PDF (needs the server; not available until
                                  staging). 2 pages: a list, a table, a page break.

Use your own phone camera photos too (see the checklist): they are the real test.
"""


def main(output: Path) -> None:
    if output.exists() and any(output.iterdir()):
        sys.exit(f"{output} is not empty; give an empty or new folder.")
    folder = output / "device-test-pack"
    folder.mkdir(parents=True)
    copies = {
        "a.pdf": "01-merge-a-2-pages.pdf",
        "b.pdf": "02-merge-b-2-pages.pdf",
        "preview-features.pdf": "03-four-pages.pdf",
        "metadata-rich.pdf": "07-metadata-rich.pdf",
        "text-scanned.pdf": "08-scanned-no-text.pdf",
        "text-partial.pdf": "09-partly-scanned.pdf",
        "encrypted.pdf": "10-password-protected.pdf",
        "corrupt.pdf": "11-damaged.pdf",
        "a.pdf#noext": "12-pdf-without-extension",
    }
    for source, target in copies.items():
        shutil.copyfile(FIXTURES / source.split("#")[0], folder / target)
    numbered_pdf(folder / "04-fifty-pages.pdf", 50, "Fifty-page test")
    numbered_pdf(folder / "05-two-hundred-pages.pdf", 200, "Two-hundred-page test")
    image_heavy_pdf(folder / "06-image-heavy-5-pages.pdf", 5)
    transparent_png(folder / "13-transparent.png")
    camera_style_jpeg(folder / "14-camera-style-rotated.jpg")
    word_document(folder / "15-word-document.docx")
    (folder / "README.txt").write_text(README, encoding="utf-8")

    archive = output / "document-platform-device-test-pack.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for path in sorted(folder.iterdir()):
            bundle.write(path, f"device-test-pack/{path.name}")
    for path in sorted(folder.iterdir()):
        print(f"{path.stat().st_size / 1024:9.1f} KB  {path.name}")
    print(f"{archive.stat().st_size / 1024:9.1f} KB  {archive.name} (zip)")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(Path(sys.argv[1]))
