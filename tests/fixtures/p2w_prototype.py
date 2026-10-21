"""Candidate B for the M5B benchmark: a first-party PDF to DOCX prototype.

Built only from permissive parts (pdfplumber / pdfminer.six: MIT, python-docx:
MIT, pypdfium2: Apache-2.0/BSD). It exists to answer one question: what quality
could we reach by building the converter ourselves? It is benchmark tooling,
not product code.

    uv run --with pdfplumber --with python-docx --with pypdfium2 --with pillow \
        python tests/fixtures/p2w_prototype.py in.pdf out.docx

What it does: layout analysis with a recursive XY-cut (reading order, columns),
paragraph and heading detection from font size, bold/italic runs, bullet and
numbered lists, bordered tables (merged across pages), embedded images, RTL
paragraphs, and image-only pages (embedded as pictures, no text invented).
What it does not do: borderless tables, footnotes, hyperlinks, form fields,
vector graphics, OCR.
"""

import io
import re
import statistics
import sys
import unicodedata
from collections import Counter

import pdfplumber
import pypdfium2 as pdfium
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Pt

ARABIC = re.compile(r"[؀-ۿݐ-ݿﭐ-﷿ﹰ-﻿]")
MARKER = re.compile(r"^(?:[•◦▪‣\-–]|\d{1,3}[.)])$")
SERIF = re.compile(r"times|serif|roman|georgia|garamond|charter|liberation ?serif", re.I)
MONO = re.compile(r"courier|mono|consolas", re.I)


class Atom:
    __slots__ = ("kind", "x0", "x1", "top", "bottom", "data")

    def __init__(self, kind, x0, x1, top, bottom, data=None):
        self.kind, self.x0, self.x1, self.top, self.bottom, self.data = kind, x0, x1, top, bottom, data


def xy_cut(atoms, gx, gy):
    """Order atoms into groups: horizontal bands first, vertical gutters when
    they are clearly columns (line tops on both sides do not line up)."""
    if len(atoms) <= 1:
        return [atoms]
    h = best_gap(atoms, "top", "bottom", gy)
    v = best_gap(atoms, "x0", "x1", gx)
    use_v = v is not None and (h is None or v[0] > h[0]) and not row_aligned(atoms, v[1])
    if v is not None and not use_v and h is None:
        return [sorted(atoms, key=lambda a: (round(a.top), a.x0))]
    cut = v if use_v else h
    if cut is None:
        return [atoms]
    _, at = cut
    key = (lambda a: a.x0) if use_v else (lambda a: a.top)
    ordered = sorted(atoms, key=key)
    first = [a for a in ordered if (a.x1 if use_v else a.bottom) <= at + 0.01]
    taken = {id(a) for a in first}
    second = [a for a in ordered if id(a) not in taken]
    if not first or not second:
        return [atoms]
    return xy_cut(first, gx, gy) + xy_cut(second, gx, gy)


def best_gap(atoms, lo, hi, minimum):
    items = sorted(atoms, key=lambda a: getattr(a, lo))
    best, reach = None, getattr(items[0], hi)
    for atom in items[1:]:
        start = getattr(atom, lo)
        gap = start - reach
        if gap >= minimum and (best is None or gap > best[0]):
            best = (gap, reach)
        reach = max(reach, getattr(atom, hi))
    return best


def row_aligned(atoms, split):
    """True when atoms either side of a vertical gutter sit on the same text
    rows (a form or table-like row layout), so rows should stay intact."""
    left = {round(a.top) for a in atoms if a.x1 <= split + 0.01 and a.kind == "word"}
    right = {round(a.top) for a in atoms if a.x0 > split and a.kind == "word"}
    if not left or not right:
        return False
    matched = sum(1 for t in left if any(abs(t - u) <= 1 for u in right))
    if matched / min(len(left), len(right)) < 0.85:
        return False
    # Two prose columns often share a baseline grid too. Their lines run almost
    # up to the gutter and there are many of them; form labels are short.
    words = [a for a in atoms if a.kind == "word" and a.x1 <= split + 0.01]
    ends = {}
    for a in words:
        ends[round(a.top)] = max(ends.get(round(a.top), 0), a.x1)
    start = min(a.x0 for a in words)
    fill = [(end - start) / max(1.0, split - start) for end in ends.values()]
    return not (len(fill) >= 6 and statistics.median(fill) >= 0.85)


def lines_of(words):
    lines = []
    for word in sorted(words, key=lambda a: (a.top, a.x0)):
        if lines and abs(lines[-1][0].top - word.top) <= max(2.0, 0.35 * (word.bottom - word.top)):
            lines[-1].append(word)
        else:
            lines.append([word])
    for line in lines:
        line.sort(key=lambda a: a.x0)
    return lines


LTR_RUN = re.compile(r"[A-Za-z0-9][A-Za-z0-9.,:/%@-]*")


def logical_arabic(text):
    """Words arrive left-to-right as drawn, in presentation forms. Reverse the
    drawn order, restore embedded digits and Latin runs, then fold the forms."""
    flipped = LTR_RUN.sub(lambda m: m.group(0)[::-1], text[::-1])
    return unicodedata.normalize("NFKC", flipped)


def is_rtl(line):
    letters = sum(len(w.data["text"]) for w in line)
    arabic = sum(len(ARABIC.findall(w.data["text"])) for w in line)
    return letters > 0 and arabic / letters > 0.5


def line_text(line):
    ordered = list(reversed(line)) if is_rtl(line) else line
    return ordered


def font_kind(name):
    if MONO.search(name or ""):
        return "Courier New"
    return "Times New Roman" if SERIF.search(name or "") else "Arial"


def runs_of(words):
    """Merge consecutive words with the same style into (text, bold, italic)."""
    runs = []
    for index, word in enumerate(words):
        name = word.data.get("fontname") or ""
        style = (bool(re.search(r"bold|black|heavy", name, re.I)), bool(re.search(r"italic|oblique", name, re.I)))
        text = word.data["text"] + (" " if index < len(words) - 1 else "")
        if runs and runs[-1][1:] == style:
            runs[-1] = (runs[-1][0] + text, *style)
        else:
            runs.append((text, *style))
    return runs


def build_paragraphs(group, body_size):
    """Turn one leaf group of word atoms into paragraph specs."""
    lines = lines_of([a for a in group if a.kind == "word"])
    paragraphs, current = [], []
    for line in lines:
        words = line_text(line)
        first = words[0].data["text"]
        marker = MARKER.match(first) and len(words) > 1
        if current and marker:
            paragraphs.append(current)
            current = []
        current.append((line, words, bool(marker)))
    if current:
        paragraphs.append(current)
    specs = []
    for para in paragraphs:
        items = []
        listed = para[0][2]
        marker_text = para[0][1][0].data["text"] if listed else ""
        for index, (_line, words, has_marker) in enumerate(para):
            used = words[1:] if has_marker and index == 0 else words
            items.extend(used)
        sizes = [w.data["size"] for w in items] or [body_size]
        size = statistics.median(sizes)
        text = join_lines([[w for w in (words[1:] if has_marker and i == 0 else words)] for i, (line, words, has_marker) in enumerate(para)])
        rtl = any(is_rtl(line) for line, _, _ in para)
        specs.append({
            "runs": runs_of_lines(para, listed),
            "text": text, "size": size, "list": marker_text if listed else None,
            "rtl": rtl, "lines": len(para),
            "bold": all(r[1] for r in runs_of_lines(para, listed)),
            "font": font_kind(items[0].data.get("fontname")) if items else "Arial",
            "left": min(l[0].x0 for l, _, _ in para),
        })
    return specs


def join_lines(lines):
    out = ""
    for words in lines:
        text = " ".join(w.data["text"] for w in words)
        if out.endswith("-") and text[:1].islower():
            out = out[:-1] + text
        else:
            out = (out + " " + text) if out else text
    return out


def runs_of_lines(para, listed):
    merged = []
    for index, (_line, words, has_marker) in enumerate(para):
        used = words[1:] if has_marker and index == 0 else words
        merged.extend(used)
    return runs_of(merged)


def add_paragraph(document, spec, heading_level, body_size):
    if spec["list"] is not None:
        style = "List Number" if re.match(r"\d", spec["list"]) else "List Bullet"
        paragraph = document.add_paragraph(style=style)
    elif heading_level:
        paragraph = document.add_paragraph(style=f"Heading {heading_level}")
    else:
        paragraph = document.add_paragraph()
    if spec["rtl"]:
        properties = paragraph._p.get_or_add_pPr()
        properties.append(OxmlElement("w:bidi"))
    for text, bold, italic in spec["runs"]:
        run = paragraph.add_run(text)
        run.bold = True if bold and not heading_level else None
        run.italic = True if italic else None
        if not heading_level:
            run.font.size = Pt(round(spec["size"] * 2) / 2)
            run.font.name = spec["font"]
        if spec["rtl"]:
            rpr = run._r.get_or_add_rPr()
            rpr.append(OxmlElement("w:rtl"))
    return paragraph


def add_table(document, rows):
    columns = max(len(r) for r in rows)
    table = document.add_table(rows=0, cols=columns)
    table.style = "Table Grid"
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = (value or "").replace("\n", " ")
            for p in cells[index].paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
    document.add_paragraph()
    return table


def emit_words(document, buffer, body):
    """Write the buffered word atoms as paragraphs and headings, then empty it."""
    for spec in build_paragraphs(list(buffer), body):
        big = spec["size"] >= body * 1.18 and spec["lines"] <= 3
        level = (1 if spec["size"] >= body * 1.6 else 2) if big else 0
        add_paragraph(document, spec, level, body)
    buffer.clear()


def picture_bytes(image):
    buffer = io.BytesIO()
    image.convert("RGB").save(buffer, "PNG")
    return buffer.getvalue()


def convert(source, target):
    pdf = pdfium.PdfDocument(source)
    document = Document()
    section = document.sections[0]
    sizes = Counter()
    pages_atoms = []
    margin = 54
    with pdfplumber.open(source) as plumber:
        for index, page in enumerate(plumber.pages):
            handle = pdf[index]
            text_page = handle.get_textpage()
            chars = text_page.count_chars()
            text_page.close()
            handle.close()
            if chars == 0:
                pages_atoms.append(("scanned", picture_bytes(page.to_image(resolution=120).original), (page.width, page.height)))
                page.flush_cache()
                continue
            tables = [t for t in page.find_tables() if len(t.rows) >= 2 and len(t.rows[0].cells) >= 2]
            boxes = [t.bbox for t in tables]
            atoms = []
            words = page.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False, use_text_flow=False)
            for w in words:
                w = {k: w[k] for k in ("text", "x0", "x1", "top", "bottom", "fontname", "size")}
                if ARABIC.search(w["text"]):
                    w["text"] = logical_arabic(w["text"])
                cx, cy = (w["x0"] + w["x1"]) / 2, (w["top"] + w["bottom"]) / 2
                if any(b[0] <= cx <= b[2] and b[1] <= cy <= b[3] for b in boxes):
                    continue
                atoms.append(Atom("word", w["x0"], w["x1"], w["top"], w["bottom"], w))
                sizes[round(w["size"])] += len(w["text"])
            for table, box in zip(tables, boxes, strict=True):
                atoms.append(Atom("table", box[0], box[2], box[1], box[3], table.extract()))
            for image in page.images:
                bbox = (image["x0"], image["top"], image["x1"], image["bottom"])
                if (bbox[2] - bbox[0]) < 8 or (bbox[3] - bbox[1]) < 8:
                    continue
                crop = picture_bytes(page.crop(bbox).to_image(resolution=150).original)
                atoms.append(Atom("image", bbox[0], bbox[2], bbox[1], bbox[3], data=crop))
            pages_atoms.append(("text", (page.width, page.height), atoms))
            page.flush_cache()
        body = sizes.most_common(1)[0][0] if sizes else 11
        first = pages_atoms[0][2] if pages_atoms and pages_atoms[0][0] == "scanned" else (pages_atoms[0][1] if pages_atoms else None)
        if first is not None:
            section.page_width, section.page_height = Pt(first[0]), Pt(first[1])
            xs = [a.x0 for kind, _, atoms in pages_atoms if kind == "text" for a in atoms]
            margin = min(max(min(xs, default=54) if xs else 54, 24), 108)
            section.left_margin = section.right_margin = Pt(margin)
            section.top_margin = section.bottom_margin = Pt(54)
        previous_table = None
        for number, (kind, page, atoms) in enumerate(pages_atoms):
            if number and kind != "skip":
                document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            if kind == "scanned":
                document.add_picture(io.BytesIO(page), width=Pt(atoms[0] - 2 * margin))
                previous_table = None
                continue
            page_width = page[0]
            groups = xy_cut(atoms, gx=max(12.0, body * 1.0), gy=max(6.0, body * 0.6)) if atoms else []
            for group in groups:
                buffer = []
                for atom in [*sorted(group, key=lambda a: (round(a.top), a.x0)), None]:
                    if atom is not None and atom.kind == "word":
                        buffer.append(atom)
                        continue
                    if buffer:  # text before a table or picture, or at the end of the group
                        previous_table = None
                        emit_words(document, buffer, body)
                    if atom is None:
                        break
                    if atom.kind == "table":
                        rows = atom.data
                        if previous_table is not None and len(rows[0]) == previous_table["columns"]:
                            for row in rows:
                                cells = previous_table["table"].add_row().cells
                                for i, value in enumerate(row[: len(cells)]):
                                    cells[i].text = (value or "").replace("\n", " ")
                        else:
                            table = add_table(document, rows)
                            previous_table = {"table": table, "columns": len(rows[0])}
                    else:
                        previous_table = None
                        document.add_picture(io.BytesIO(atom.data), width=Pt(min(atom.x1 - atom.x0, page_width - 2 * margin)))
    document.save(target)


if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
