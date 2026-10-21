"""Synthetic corpus for the M5A Compress PDF benchmark.

Writes ignored data to .tools/compress-benchmarks/{corpus,stress}. Every file is
authored here from public-domain sample images (scikit-image) and generated
text; no user documents are involved. The corpus is deterministic.

Run (pins nothing in the repo; dependencies are installed on the fly):

    uv run --with pikepdf --with pillow --with numpy --with scikit-image \
        --with fonttools --with python-docx \
        python tests/fixtures/generate_compress_benchmark.py [--skip-soffice]

Classes
  A  vector/text documents (structure only, no images)
  B  photo documents
  C  scans (colour JPEG, grey Flate, bilevel Flate)
  D  screenshots and flat graphics
  E  LibreOffice exports (the same producer as the Word to PDF tool)
  G  stress: decode-size, page-count and upload-size boundaries
"""

import argparse
import io
import random
import shutil
import subprocess
import tempfile
import zlib
from pathlib import Path

import numpy as np
import pikepdf
from fontTools.ttLib import TTFont
from pikepdf import Array, Dictionary, Name, Stream
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from skimage import data

ROOT = Path(__file__).resolve().parents[2] / ".tools" / "compress-benchmarks"
A4 = (595.0, 842.0)
SERIF = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
SANS = "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"
BASES = ["rocket", "coffee", "astronaut", "chelsea"]
WORDS = (
    "the of and to in is that for it as with was on are by this be from or have an "
    "they which one you were all we can her has there been if more when will would "
    "who so no document platform privacy server conversion report quarter revenue "
    "customer project schedule budget invoice agreement section policy review"
).split()
RAW = dict(
    compress_streams=False,
    object_stream_mode=pikepdf.ObjectStreamMode.disable,
    stream_decode_level=pikepdf.StreamDecodeLevel.none,
    recompress_flate=False,
)


def lorem(rng, count):
    return " ".join(rng.choice(WORDS) for _ in range(count))


def photo(width, height, seed=1, sigma=9.0, base="rocket"):
    """A photograph-like image: a real photo resampled, plus multi-scale noise."""
    source = Image.fromarray(getattr(data, base)())
    sw, sh = source.size
    ratio = width / height
    if sw / sh > ratio:
        nw = int(sh * ratio)
        source = source.crop(((sw - nw) // 2, 0, (sw - nw) // 2 + nw, sh))
    else:
        nh = int(sw / ratio)
        source = source.crop((0, (sh - nh) // 2, sw, (sh - nh) // 2 + nh))
    array = np.asarray(source.resize((width, height), Image.LANCZOS)).astype(np.int16)
    rng = np.random.default_rng(seed)
    for scale, amplitude in ((1, sigma), (3, sigma * 0.8), (9, sigma * 0.6), (27, sigma * 0.4)):
        small = rng.normal(0, amplitude, (height // scale + 2, width // scale + 2)).astype(np.float32)
        layer = Image.fromarray(small, mode="F").resize((width, height), Image.BILINEAR)
        array += np.asarray(layer)[..., None].astype(np.int16)
    return Image.fromarray(np.clip(array, 0, 255).astype(np.uint8))


def scan_page(seed, color=False, dpi=300):
    """An A4 scanned text page: rendered text, skew, blur, vignette and paper noise."""
    rng = random.Random(seed)
    width, height = int(8.27 * dpi), int(11.69 * dpi)
    image = Image.new("L", (width, height), 245)
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype(SERIF, int(11 * dpi / 72))
    y, leading = int(1.0 * dpi), int(14.5 * dpi / 72)
    while y < height - int(1.0 * dpi):
        if rng.random() < 0.12:
            y += leading
            continue
        draw.text((int(1.0 * dpi), y), lorem(rng, rng.randint(11, 15)), font=font, fill=22)
        y += leading
    image = image.rotate(rng.uniform(-0.4, 0.4), resample=Image.BICUBIC, fillcolor=240)
    image = image.filter(ImageFilter.GaussianBlur(0.7))
    array = np.asarray(image).astype(np.int16)
    noise = np.random.default_rng(seed)
    yy = np.linspace(-1, 1, height)[:, None]
    xx = np.linspace(-1, 1, width)[None, :]
    array = array - (12 * (xx**2 + yy**2)).astype(np.int16)
    array = array + noise.normal(0, 4, (height, width)).astype(np.int16)
    grey = Image.fromarray(np.clip(array, 0, 255).astype(np.uint8))
    if color:
        rgb = np.stack([np.asarray(grey).astype(np.int16) + o for o in (4, 0, -9)], -1)
        return Image.fromarray(np.clip(rgb, 0, 255).astype(np.uint8))
    return grey


def screenshot(seed, width=1920, height=1080):
    rng = random.Random(seed)
    image = Image.new("RGB", (width, height), (250, 250, 252))
    draw = ImageDraw.Draw(image)
    font, bold = ImageFont.truetype(SANS, 22), ImageFont.truetype(SANS, 34)
    draw.rectangle([0, 0, width, 64], fill=(32, 64, 128))
    draw.text((24, 16), "Dashboard", font=bold, fill=(255, 255, 255))
    for i in range(9):
        y = 100 + i * 96
        draw.rounded_rectangle([40, y, width - 40, y + 80], 10, outline=(200, 205, 215), fill=(255, 255, 255))
        draw.text((64, y + 24), lorem(rng, 12), font=font, fill=(40, 40, 50))
        draw.rectangle([width - 360, y + 20, width - 80, y + 60], fill=(40 + i * 18, 160, 100))
    return image


def image_xobject(pdf, image, kind, quality=85):
    if kind == "jpeg":
        buffer = io.BytesIO()
        image.save(buffer, "JPEG", quality=quality)
        payload, flt = buffer.getvalue(), Name.DCTDecode
    elif kind == "flate":
        payload, flt = zlib.compress(image.tobytes(), 6), Name.FlateDecode
    else:
        bilevel = image.convert("L").point(lambda v: 255 if v > 128 else 0).convert("1")
        payload, flt = zlib.compress(bilevel.tobytes(), 9), Name.FlateDecode
    stream = Stream(pdf, payload)
    stream.Type, stream.Subtype = Name.XObject, Name.Image
    stream.Width, stream.Height, stream.Filter = image.width, image.height, flt
    if kind == "bilevel":
        stream.ColorSpace, stream.BitsPerComponent = Name.DeviceGray, 1
    else:
        stream.ColorSpace = Name.DeviceGray if image.mode == "L" else Name.DeviceRGB
        stream.BitsPerComponent = 8
    return pdf.make_indirect(stream)


def text_ops(rng, lines, y0=800):
    ops = [f"BT /F1 10 Tf 13 TL 56 {y0} Td"]
    for _ in range(lines):
        text = lorem(rng, rng.randint(9, 14))
        ops.append(f"({text}) '")
    ops.append("ET")
    return "\n".join(ops) + "\n"


def embed_ttf(pdf, path):
    raw = Path(path).read_bytes()
    font = TTFont(path)
    cmap, hmtx, upm = font.getBestCmap(), font["hmtx"], font["head"].unitsPerEm
    widths = [int(hmtx[cmap[c]][0] * 1000 / upm) if c in cmap else 0 for c in range(32, 256)]
    file = Stream(pdf, raw)
    file.Length1 = len(raw)
    name = Name("/" + font["name"].getDebugName(6))
    descriptor = pdf.make_indirect(Dictionary(
        Type=Name.FontDescriptor, FontName=name, Flags=32, FontBBox=Array([-200, -300, 1200, 1000]),
        ItalicAngle=0, Ascent=900, Descent=-200, CapHeight=700, StemV=80, FontFile2=pdf.make_indirect(file)))
    return pdf.make_indirect(Dictionary(
        Type=Name.Font, Subtype=Name.TrueType, BaseFont=name, FirstChar=32, LastChar=255,
        Widths=Array(widths), FontDescriptor=descriptor, Encoding=Name.WinAnsiEncoding))


def build(path, pages, *, font="std", save=None):
    """pages: [{text: lines, images: [(image, kind, quality, x, y, width_pt)]}]"""
    pdf = pikepdf.Pdf.new()
    if font == "ttf":
        f1 = embed_ttf(pdf, SANS)
    else:
        f1 = pdf.make_indirect(Dictionary(
            Type=Name.Font, Subtype=Name.Type1, BaseFont=Name.Helvetica, Encoding=Name.WinAnsiEncoding))
    rng = random.Random(7)
    for spec in pages:
        xobjects, content = {}, ""
        if spec.get("text"):
            content += text_ops(rng, spec["text"], spec.get("text_y", 800))
        for i, (image, kind, quality, x, y, width_pt) in enumerate(spec.get("images", [])):
            height_pt = width_pt * image.height / image.width
            xobjects[f"/Im{i}"] = image_xobject(pdf, image, kind, quality)
            content += f"q {width_pt:.2f} 0 0 {height_pt:.2f} {x:.2f} {y:.2f} cm /Im{i} Do Q\n"
        resources = Dictionary(Font=Dictionary(F1=f1))
        if xobjects:
            resources.XObject = Dictionary(xobjects)
        stream = pdf.make_indirect(Stream(pdf, content.encode("latin-1")))
        page = pdf.make_indirect(Dictionary(
            Type=Name.Page, MediaBox=Array([0, 0, *A4]), Resources=resources, Contents=stream))
        pdf.pages.append(pikepdf.Page(page))
    pdf.save(path, **(save or RAW))
    return path


def libreoffice_documents(out, workdir):
    """Word documents converted by LibreOffice, the same producer as Word to PDF."""
    import docx
    from docx.shared import Inches

    def make(name, builder):
        document = docx.Document()
        builder(document)
        source = workdir / f"{name}.docx"
        document.save(source)
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(workdir), str(source)],
            check=True, capture_output=True, timeout=300)
        shutil.copy(workdir / f"{name}.pdf", out / f"{name}.pdf")

    def picture(document, image, width_in, quality=90):
        buffer = io.BytesIO()
        image.save(buffer, "JPEG", quality=quality)
        buffer.seek(0)
        document.add_picture(buffer, width=Inches(width_in))

    rng = random.Random(11)

    def report(document):
        for i in range(8):
            document.add_heading(f"Section {i + 1}", level=1)
            for _ in range(6):
                document.add_paragraph(lorem(rng, 90))
            picture(document, photo(2400, 1600, seed=i, base=BASES[i % 4]), 6.0)
            document.add_paragraph(lorem(rng, 12))

    def long_document(document):
        for i in range(120):
            document.add_heading(f"Chapter {i + 1}", level=2)
            for _ in range(8):
                document.add_paragraph(lorem(rng, 110))
            if i % 4 == 0:
                buffer = io.BytesIO()
                screenshot(i).save(buffer, "PNG")
                buffer.seek(0)
                document.add_picture(buffer, width=Inches(6.0))

    def text_only(document):
        for i in range(10):
            document.add_heading(f"Part {i + 1}", level=1)
            for _ in range(10):
                document.add_paragraph(lorem(rng, 120))

    make("E01-lo-report-photos", report)
    make("E02-lo-long-doc-images", long_document)
    make("E03-lo-text-only", text_only)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--skip-soffice", action="store_true")
    parser.add_argument("--skip-stress", action="store_true")
    args = parser.parse_args()
    corpus, stress = ROOT / "corpus", ROOT / "stress"
    corpus.mkdir(parents=True, exist_ok=True)
    stress.mkdir(parents=True, exist_ok=True)

    def report(name, path):
        print(f"{name:34s} {path.stat().st_size / 1e6:8.2f} MB", flush=True)

    pages = [dict(text=62) for _ in range(40)]
    report("A01 text, unoptimized", build(corpus / "A01-text-unoptimized.pdf", pages))
    report("A02 text, already optimized", build(
        corpus / "A02-text-optimized.pdf", pages,
        save=dict(compress_streams=True, object_stream_mode=pikepdf.ObjectStreamMode.generate, recompress_flate=True)))
    report("A03 embedded full TTF, raw", build(corpus / "A03-embedded-font-raw.pdf", pages, font="ttf"))

    specs = []
    for i in range(6):
        images = [(photo(3000, 2000, seed=i, base=BASES[i % 4]), "jpeg", 90, 56, 430, 396)] if i < 5 else []
        specs.append(dict(text=24, images=images))
    report("B01 office, 5 hi-res photos", build(corpus / "B01-office-photos.pdf", specs))
    specs = [dict(images=[(photo(4000, 3000, seed=10 + i, base=BASES[i % 4]), "jpeg", 92, 20, 300, 555)]) for i in range(4)]
    report("B02 camera photos 4 x 12 MP", build(corpus / "B02-camera-photos.pdf", specs))

    colour = [scan_page(100 + i, color=True) for i in range(8)]
    report("C01 colour scan JPEG q85 x8", build(
        corpus / "C01-scan-color-jpeg.pdf", [dict(images=[(im, "jpeg", 85, 0, 0, 595.0)]) for im in colour]))
    grey = [scan_page(200 + i) for i in range(3)]
    report("C02 grey scan Flate x3", build(
        corpus / "C02-scan-gray-flate.pdf", [dict(images=[(im, "flate", 0, 0, 0, 595.0)]) for im in grey]))
    report("C03 bilevel scan Flate x3", build(
        corpus / "C03-scan-bilevel-flate.pdf", [dict(images=[(im, "bilevel", 0, 0, 0, 595.0)]) for im in grey]))

    specs = [dict(text=6, text_y=820, images=[(screenshot(300 + i), "flate", 0, 40, 100, 515)]) for i in range(10)]
    report("D01 screenshots Flate x10", build(corpus / "D01-screenshots-flate.pdf", specs))

    if not args.skip_soffice and shutil.which("soffice"):
        with tempfile.TemporaryDirectory() as workdir:
            libreoffice_documents(corpus, Path(workdir))
        for name in ("E01-lo-report-photos", "E02-lo-long-doc-images", "E03-lo-text-only"):
            report(name, corpus / f"{name}.pdf")

    if args.skip_stress:
        return
    # G01: 100 MP JPEG (smooth, so the file is small but decoding is huge).
    width, height = 12500, 8000
    array = np.zeros((height, width, 3), np.uint8)
    array[:, :, 0] = np.linspace(0, 255, width, dtype=np.uint8)[None, :]
    array[:, :, 1] = np.linspace(0, 255, height, dtype=np.uint8)[:, None]
    array[:, :, 2] = 90
    report("G01 100 MP JPEG", build(
        stress / "G01-jpeg-100MP.pdf", [dict(images=[(Image.fromarray(array), "jpeg", 80, 20, 300, 555)])]))
    # G02: 40 MP Flate RGB.
    width, height = 8000, 5000
    array = np.zeros((height, width, 3), np.uint8)
    array[:, :, 0] = np.linspace(0, 255, width, dtype=np.uint8)[None, :]
    array[:, :, 1] = np.linspace(0, 255, height, dtype=np.uint8)[:, None]
    array[:, :, 2] = 40
    report("G02 40 MP Flate RGB", build(
        stress / "G02-flate-40MP.pdf", [dict(images=[(Image.fromarray(array), "flate", 0, 20, 300, 555)])]))
    # G03/G05: page-count boundaries for the structure pass.
    report("G03 400 pages text", build(stress / "G03-400pages-text.pdf", [dict(text=60) for _ in range(400)]))
    report("G05 1000 pages text", build(stress / "G05-1000pages-text.pdf", [dict(text=40) for _ in range(1000)]))
    # G04: 150 pages with one 1200x900 JPEG each (about 37 MB).
    specs = [dict(text=4, text_y=820, images=[(photo(1200, 900, seed=i % 7, base=BASES[i % 4]), "jpeg", 85, 30, 300, 535)]) for i in range(150)]
    report("G04 150 pages x image", build(stress / "G04-150pages-images.pdf", specs))
    # G06: 8000 x 8000 (64 MP) JPEG, exactly the planned per-image decode cap.
    width = height = 8000
    array = np.zeros((height, width, 3), np.uint8)
    rng = np.random.default_rng(5)
    array[:, :, 0] = np.linspace(0, 255, width, dtype=np.uint8)[None, :]
    array[:, :, 1] = np.linspace(0, 255, height, dtype=np.uint8)[:, None]
    array[:, :, 2] = rng.integers(60, 120, (height, width), dtype=np.uint8)
    report("G06 64 MP JPEG", build(
        stress / "G06-jpeg-64MP.pdf", [dict(images=[(Image.fromarray(array), "jpeg", 85, 20, 300, 555)])]))


if __name__ == "__main__":
    main()
